"""Build V3 research figures and editable reports from verified experiment artifacts.

No experiment score is embedded in this module. V2 artifacts are read-only references.
"""
import argparse
import hashlib
import json
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

TITLE = '电循智策——基于BMS多维运行数据的动力电池健康评估与循环利用决策研究'
COLORS = ['#0072B2', '#D55E00', '#009E73', '#CC79A7']
LABELS = {'train_mean': '训练均值基线', 'ridge': 'Ridge', 'random_forest': '随机森林',
          'without_soc_current': '去 SOC 与电流', 'without_temperature': '去温度',
          'without_cell_gap': '去单体压差', 'noise_1pct': '1% 特征噪声',
          'noise_5pct': '5% 特征噪声', 'missing_10pct': '10% 特征缺失'}
NAMES = ['01_dataset_split', '02_model_mae', '03_rf_predictions', '04_vehicle_mae',
         '05_random_seeds', '06_group_cv', '07_feature_ablation', '08_input_stress',
         '09_absolute_errors', '10_quality_control']
CAPTIONS = ['三集合的独立车辆数与充电片段数', '固定测试集三模型容量代理 MAE',
            '固定测试集随机森林预测值与容量代理标签', '固定测试集随机森林逐车 MAE',
            '三个随机种子的随机森林测试 MAE', '训练车辆内五折分组验证 MAE',
            '随机森林特征消融结果', '冻结随机森林的特征层模拟压力测试',
            '固定测试集三模型绝对误差经验分布', '实验与图表质量控制状态']
CHECKS = [('split', '车辆划分'), ('fields', '输入字段'), ('recompute', '指标复算'),
          ('reload', '模型重载'), ('predictions', '逐条预测'), ('figures', '图表与源数据'),
          ('seeds', '随机种子'), ('group_cv', '分组交叉验证'), ('noise', '模拟噪声'),
          ('missing', '模拟缺失')]


def load_inputs(result, source):
    metrics = pd.read_csv(result / 'metrics.csv')
    predictions = pd.read_csv(result / 'predictions.csv.gz')
    cv = pd.read_csv(result / 'group_cv.csv')
    vehicle = pd.read_csv(result / 'per_vehicle.csv')
    features = pd.read_csv(source)
    for frame, cols in [(metrics, ['mae_ah', 'rmse_ah', 'r2']),
                        (predictions, ['proxy_capacity_ah', 'predicted_ah']), (cv, ['mae_ah'])]:
        if not np.isfinite(frame[cols]).all().all():
            raise ValueError('Non-finite report input')
    return metrics, predictions, cv, vehicle, features


def select(frame, names):
    selected = frame.loc[(frame.split == 'test') & frame.experiment.isin(names)].copy()
    selected['order'] = selected.experiment.map({name: i for i, name in enumerate(names)})
    selected = selected.sort_values('order').drop(columns='order')
    if len(selected) != len(names):
        raise ValueError('Expected one test result per experiment: ' + str(names))
    return selected


def style():
    plt.rcParams.update({'font.sans-serif': ['Microsoft YaHei', 'SimHei', 'DejaVu Sans'],
        'font.size': 9, 'axes.labelsize': 10, 'xtick.labelsize': 9, 'ytick.labelsize': 9,
        'axes.unicode_minus': False, 'axes.spines.top': False, 'axes.spines.right': False,
        'axes.edgecolor': '#555555', 'axes.linewidth': .7, 'figure.facecolor': 'white',
        'axes.facecolor': 'white', 'savefig.facecolor': 'white', 'svg.fonttype': 'path'})


def canvas(ncols=1):
    return plt.subplots(1, ncols, figsize=(6.4, 3.6), layout='constrained')


def dots(ax, frame, labels=None):
    y = np.arange(len(frame))
    ax.scatter(frame.mae_ah, y, s=42, color=COLORS[0], zorder=3)
    ax.hlines(y, 0, frame.mae_ah, color='#B5C5CD', linewidth=1)
    for value, position in zip(frame.mae_ah, y):
        ax.annotate(f'{value:.3f}', (value, position), xytext=(6, 0), textcoords='offset points', va='center', fontsize=9)
    ax.set_yticks(y, labels or [LABELS[x] for x in frame.experiment])
    ax.invert_yaxis()
    ax.set_xlim(0, frame.mae_ah.max() * 1.23)
    ax.set_ylim(len(frame) - .45, -.55)
    ax.set_xlabel('容量代理 MAE / Ah（越低越好）')
    ax.grid(axis='x', color='#E5E5E5', linewidth=.6)
    ax.set_axisbelow(True)


def build_figures(result, source, output, verification):
    style()
    metrics, pred, cv, vehicle, features = load_inputs(result, source)
    dirs = {key: output / folder for key, folder in [('png', '05_figures_png'), ('svg', '06_figures_svg'), ('csv', '07_figure_source_data')]}
    for directory in dirs.values(): directory.mkdir(parents=True, exist_ok=True)
    audit = []

    def save(index, fig, data):
        name = NAMES[index - 1]
        for extension in ['png', 'svg']:
            fig.savefig(dirs[extension] / f'{name}.{extension}', dpi=300)
        data.to_csv(dirs['csv'] / f'{name}.csv', index=False, encoding='utf-8-sig')
        plt.close(fig)
        audit.append({'figure': name, 'source_rows': len(data), 'width_px': 1920, 'height_px': 1080,
                      'dpi': 300, 'source_sha256': hashlib.sha256((dirs['csv'] / f'{name}.csv').read_bytes()).hexdigest()})

    split = features.groupby('split').agg(vehicles=('vehicle', 'nunique'), sessions=('vehicle', 'size')).reindex(['train', 'validation', 'test']).reset_index()
    fig, axes = canvas(2)
    for ax, field, xlabel in zip(axes, ['vehicles', 'sessions'], ['独立车辆数 / 辆', '充电片段数 / 段']):
        ax.hlines(range(3), 0, split[field], color='#B5C5CD', linewidth=1)
        ax.scatter(split[field], range(3), s=40, color=COLORS[:3], zorder=3)
        for i, value in enumerate(split[field]):
            ax.annotate(f'{value:,}', (value, i), xytext=(5, 0), textcoords='offset points', va='center', fontsize=9)
        ax.set_yticks(range(3), ['训练', '验证', '测试'])
        ax.set_ylim(2.5, -.5)
        ax.set_xlim(0, split[field].max() * 1.3)
        ax.set_xlabel(xlabel)
        ax.xaxis.grid(True, color='#E5E5E5', linewidth=.6)
        ax.set_axisbelow(True)
    save(1, fig, split)

    base = select(metrics, ['train_mean', 'ridge', 'random_forest'])
    rf = base[base.experiment == 'random_forest'].iloc[0]
    fig, ax = canvas(); dots(ax, base); save(2, fig, base)

    p = pred[(pred.split == 'test') & (pred.experiment == 'random_forest') & (pred.seed == 42)].copy()
    if len(p) != rf.n: raise ValueError('Prediction/metric row-count mismatch')
    fig, ax = canvas()
    for i, (vid, group) in enumerate(p.groupby('vehicle')):
        ax.scatter(group.proxy_capacity_ah, group.predicted_ah, s=4, alpha=.34, linewidths=0,
                   color=COLORS[i % len(COLORS)], label=f'车辆 {vid}')
    bounds = [min(p.proxy_capacity_ah.min(), p.predicted_ah.min()), max(p.proxy_capacity_ah.max(), p.predicted_ah.max())]
    padding = (bounds[1] - bounds[0]) * .04
    lim = [bounds[0] - padding, bounds[1] + padding]
    ax.plot(lim, lim, '--', color='#444444', linewidth=.8, label='一致线')
    ax.set(xlim=lim, ylim=lim, xlabel='电流积分 / SOC 容量代理标签 / Ah', ylabel='随机森林预测 / Ah')
    ax.text(.03, .96, f'MAE = {rf.mae_ah:.3f} Ah\nRMSE = {rf.rmse_ah:.3f} Ah\nR² = {rf.r2:.3f}', transform=ax.transAxes, va='top', fontsize=9)
    ax.legend(loc='lower right', fontsize=7, frameon=False, markerscale=2)
    save(3, fig, p)

    per = vehicle[(vehicle.split == 'test') & (vehicle.experiment == 'random_forest') & (vehicle.seed == 42)].sort_values('vehicle')
    fig, ax = canvas(); dots(ax, per, [f'车辆 {int(r.vehicle)}（n={int(r.sessions)}）' for _, r in per.iterrows()]); save(4, fig, per)

    seeds = metrics[(metrics.split == 'test') & metrics.experiment.isin(['random_forest', 'rf_seed'])].sort_values('seed').copy()
    if seeds.seed.tolist() != [17, 42, 2026]: raise ValueError('Seeds must be exactly 17, 42, 2026')
    mean, sd = seeds.mae_ah.mean(), seeds.mae_ah.std(ddof=1)
    seeds['mean_mae_ah'] = mean; seeds['sample_sd_mae_ah'] = sd
    fig, ax = canvas()
    ax.scatter(range(3), seeds.mae_ah, color=COLORS[0], s=45, label='单次训练')
    ax.axhline(mean, color=COLORS[1], linewidth=1, label='三次均值')
    ax.axhspan(mean - sd, mean + sd, color=COLORS[1], alpha=.13, label='均值 ± 样本 SD')
    ax.text(.03, .94, f'均值 = {mean:.4f} Ah；样本 SD = {sd:.4f} Ah', transform=ax.transAxes, va='top', fontsize=9)
    for i, value in enumerate(seeds.mae_ah): ax.annotate(f'{value:.4f}', (i, value), xytext=(0, 7), textcoords='offset points', ha='center')
    ax.set_xticks(range(3), seeds.seed.astype(str)); ax.set_xlabel('随机种子'); ax.set_ylabel('容量代理 MAE / Ah')
    span = max(float(np.ptp(seeds.mae_ah)), .01)
    ax.set_ylim(mean - span * 1.5, mean + span * 1.6)
    ax.set_xlim(-.4, 2.4); ax.legend(frameon=False, fontsize=8, loc='lower right')
    save(5, fig, seeds)

    fig, ax = canvas()
    cv_data = cv.copy()
    for i, model in enumerate(['train_mean', 'ridge', 'random_forest']):
        group = cv[cv.model == model].sort_values('fold')
        if group.fold.tolist() != [1, 2, 3, 4, 5]: raise ValueError('Missing CV fold')
        average, deviation = group.mae_ah.mean(), group.mae_ah.std(ddof=1)
        ax.scatter(np.arange(5) * .065 + i - .13, group.mae_ah, s=20, color=COLORS[i], alpha=.65)
        ax.errorbar(i + .24, average, yerr=deviation, fmt='D', color=COLORS[i], capsize=4, markersize=5)
        ax.annotate(f'{average:.3f} ± {deviation:.3f}', (i, group.mae_ah.max()), xytext=(0, 12), textcoords='offset points', ha='center', fontsize=8)
        cv_data.loc[cv_data.model == model, 'fold_mean_mae_ah'] = average
        cv_data.loc[cv_data.model == model, 'fold_sample_sd_mae_ah'] = deviation
    ax.set_xticks(range(3), [LABELS[k] for k in ['train_mean', 'ridge', 'random_forest']])
    ax.set_ylabel('容量代理 MAE / Ah'); ax.set_ylim(0, cv.mae_ah.max() * 1.22)
    ax.text(.02, .02, '圆点：每折；菱形与误差线：五折均值 ± 样本 SD', transform=ax.transAxes, fontsize=8)
    ax.yaxis.grid(True, color='#E5E5E5', linewidth=.6); ax.set_axisbelow(True)
    save(6, fig, cv_data)

    ablation = select(metrics, ['random_forest', 'without_soc_current', 'without_temperature', 'without_cell_gap'])
    fig, ax = canvas(); dots(ax, ablation, ['全特征', '去 SOC 与电流', '去温度', '去单体压差']); save(7, fig, ablation)
    stress = select(metrics, ['random_forest', 'noise_1pct', 'noise_5pct', 'missing_10pct'])
    fig, ax = canvas(); dots(ax, stress, ['原始特征', '1% 特征噪声', '5% 特征噪声', '10% 特征缺失']); save(8, fig, stress)

    distribution = pred[(pred.split == 'test') & pred.experiment.isin(['train_mean', 'ridge', 'random_forest']) & (pred.seed == 42)].copy()
    distribution['absolute_error_ah'] = abs(distribution.predicted_ah - distribution.proxy_capacity_ah)
    fig, ax = canvas()
    for i, model in enumerate(['train_mean', 'ridge', 'random_forest']):
        errors = np.sort(distribution.loc[distribution.experiment == model, 'absolute_error_ah'].to_numpy())
        ax.plot(errors, np.arange(1, len(errors) + 1) / len(errors), color=COLORS[i], linewidth=1.4, label=LABELS[model])
    ax.set(xlabel='绝对误差 / Ah', ylabel='经验累计比例', ylim=(0, 1.02), xlim=(0, distribution.absolute_error_ah.max() * 1.02))
    ax.legend(frameon=False); ax.grid(color='#E5E5E5', linewidth=.6); save(9, fig, distribution)

    verified = json.loads(verification.read_text(encoding='utf-8'))
    lookup = {c['id']: c for c in verified['checks']}
    quality = []
    for key, label in CHECKS:
        item = lookup.get(key, {})
        passed = item.get('passed') is True
        detail = str(item.get('detail', '缺少检查记录'))
        quality.append({'check': key, 'label': label, 'passed': passed, 'status': '通过' if passed else '未通过', 'detail': detail})
    quality = pd.DataFrame(quality)
    fig, ax = canvas(); ax.axis('off')
    cell = [[r.label, r.status] for _, r in quality.iterrows()]
    tab = ax.table(cellText=cell, colLabels=['质量控制项目', '状态'], colWidths=[.66, .24], bbox=[.12, .02, .76, .96], cellLoc='left')
    tab.auto_set_font_size(False); tab.set_fontsize(9)
    for (row, col), value in tab.get_celld().items():
        value.set_edgecolor('#D9D9D9'); value.set_linewidth(.6)
        if row == 0: value.set_facecolor('#EDF1F4'); value.set_text_props(weight='bold')
        if row > 0 and col == 1: value.set_text_props(color=COLORS[2] if quality.iloc[row-1].passed else COLORS[1])
    save(10, fig, quality)
    report_dir = output / '09_reports'; report_dir.mkdir(parents=True, exist_ok=True)
    (report_dir / 'figure_manifest.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    return dict(metrics=metrics, rf=rf, base=base, split=split, seeds=seeds, cv=cv, ablation=ablation, stress=stress, quality=quality, predictions=p)


def interpretations(data):
    rf, base = data['rf'], data['base']
    baseline = base[base.experiment == 'train_mean'].iloc[0]
    improvement = 100 * (1 - rf.mae_ah / baseline.mae_ah)
    seed = data['seeds']
    ab = data['ablation'].set_index('experiment').mae_ah
    stress = data['stress'].set_index('experiment').mae_ah
    return [
        '训练、验证和测试按整车隔离；同车片段只属于一个集合。片段数量不等于独立车辆样本量。',
        f'随机森林测试 MAE 为 {rf.mae_ah:.4f} Ah，相比训练均值基线降低 {improvement:.2f}%。该比较衡量对容量代理标签的拟合，不能解释为 SOH 准确率。',
        f'测试集含 {int(rf.n):,} 段、{int(rf.vehicles)} 辆车。随机森林 RMSE 为 {rf.rmse_ah:.4f} Ah，R² 为 {rf.r2:.4f}。散点颜色表示车辆，不代表新的独立测试批次。',
        '逐车 MAE 反映同一固定测试集内部的差异。总体片段加权 MAE 与四车等权 MAE 的权重不同，不应混用。',
        f'种子 17、42、2026 的 MAE 均值为 {seed.mae_ah.mean():.4f} Ah，样本标准差为 {seed.mae_ah.std(ddof=1):.4f} Ah。该波动仅反映同一数据划分上的训练随机性，不是跨车辆泛化置信区间。',
        '五折分组验证仅使用训练车 1—14，每折按整车留出，固定参数。误差线为五折样本标准差，不是标准误或置信区间，也不与固定测试集视为同一总体。',
        f'每项消融均重新训练。全特征 MAE 为 {ab["random_forest"]:.4f} Ah，去单体压差为 {ab["without_cell_gap"]:.4f} Ah；结果只支持本数据与参数下的比较，不证明压差没有物理价值或标签独立有效。',
        f'冻结模型与标签，仅扰动特征。1%/5% 噪声 MAE 分别为 {stress["noise_1pct"]:.4f}/{stress["noise_5pct"]:.4f} Ah，10% 缺失为 {stress["missing_10pct"]:.4f} Ah。该模拟不等于实际 BMS 传感器误差试验。',
        '经验累计曲线展示各模型从小误差到尾部误差的完整分布；同一横坐标处累计比例越高，表示该阈值内片段比例越高。连续片段相关，不能按片段数量推断独立证据强度。',
        '状态来自本轮验证记录；通过表示对应计算或产物一致性检查通过，不表示真实电池健康、安全或退役处置有效性通过。'
    ]


def new_doc():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(.72); section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.82); section.right_margin = Inches(.82)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
        s = doc.styles[name]
        s.font.name = 'Calibri'; s._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.font.size = Pt(11 if name == 'Normal' else 10 if name == 'Caption' else 16 if name == 'Title' else 13)
        s.paragraph_format.space_after = Pt(7)
        for border in s.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.18
    doc.styles['Title'].font.bold = True
    doc.styles['Subtitle'].font.italic = False
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
    footer = section.footer.paragraphs[0]; footer.alignment = 2
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = '公开实车充电数据容量代理研究证据 V3'
    return doc


def paragraph(doc, text):
    return doc.add_paragraph(text)


def table_doc(doc, title, headers, rows):
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p = doc.add_paragraph(title, 'Caption'); p.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.autofit = False
    widths = [1.5, 5.2] if len(headers) == 2 else [2.2] + [1.5] * (len(headers)-1)
    for column, width in zip(table.columns, widths): column.width = Inches(width)
    for cell, value in zip(table.rows[0].cells, headers): cell.text = str(value)
    repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, value in zip(table.add_row().cells, row): cell.text = str(value)
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            props.append(borders)
            if i == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EEF3'); props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(5); p.paragraph_format.space_before = Pt(5)
                for run in p.runs: run.font.size = Pt(10); run.bold = i == 0
        for cell, width in zip(row.cells, widths): cell.width = Inches(width)
    return table


def figure_doc(doc, output, number, explanation, width=6.35, display_number=None):
    from docx.shared import Inches, Pt
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1); p.paragraph_format.keep_with_next = True
    p.alignment = 1
    p.add_run().add_picture(str(output / '05_figures_png' / (NAMES[number-1] + '.png')), width=Inches(width))
    caption = doc.add_paragraph(f'图 {display_number or number}  {CAPTIONS[number-1]}', 'Caption')
    caption.alignment = 1; caption.paragraph_format.keep_with_next = True
    p = paragraph(doc, explanation); p.paragraph_format.space_after = Pt(5)


def references(doc):
    doc.add_heading('参考文献与可追溯产物', 1)
    for text in [
        '[1] Deng Z, Xu L, Liu H, Hu X, Duan Z, Xu Y. Prognostics of battery capacity based on charging data and data-driven methods for on-road vehicles. Applied Energy, 2023, 339: 120954. DOI: 10.1016/j.apenergy.2023.120954.',
        '[2] BatICM. battery-charging-data-of-on-road-electric-vehicles. https://github.com/BatICM/battery-charging-data-of-on-road-electric-vehicles 。原作者数据说明：BAIC EU500、CATL NCM、145 Ah。',
        '[3] shiyunliu-battery. battery-charging-data-of-on-road-electric-vehicles. https://github.com/shiyunliu-battery/battery-charging-data-of-on-road-electric-vehicles 。下载版本 904a336bc4a8de05acdec2598708fd787bbdb8e3；来源许可与 SHA-256 链随研究产物保留。',
        '本轮数值直接读取 metrics.csv、predictions.csv.gz、group_cv.csv、per_vehicle.csv；图表逐项源数据见 07_figure_source_data。划分、模型重载及指标复算记录见 08_tests_and_verification/verification.json。'
    ]: paragraph(doc, text)


def build_documents(data, output):
    if not data['quality'].passed.all(): raise ValueError('Formal documents require all recorded checks to pass')
    docs = output / '10_docx'; docs.mkdir(parents=True, exist_ok=True)
    notes = interpretations(data)
    rf = data['rf']
    doc = new_doc()
    doc.add_paragraph(TITLE, 'Title')
    doc.add_paragraph('挑战杯研究证据报告  最终复核版 V3', 'Subtitle')
    doc.add_heading('研究结论与证据范围', 1)
    paragraph(doc, f'本报告面向研究论证与申报材料复核，整理公开实车充电数据上的容量代理建模、车辆隔离评估及稳健性补充分析。固定测试集随机森林 MAE 为 {rf.mae_ah:.4f} Ah，RMSE 为 {rf.rmse_ah:.4f} Ah，R² 为 {rf.r2:.4f}。这些数值衡量模型预测与电流积分及 SOC 变化推算标签的一致程度，不是实测 SOH 精度。')
    paragraph(doc, '研究已经形成可追溯的特征缓存、三模型对比、消融、训练车辆内分组验证、随机种子重复训练和特征层模拟压力测试。实际 SOH、RUL、故障预警与退役决策仍缺少独立实测标签和结局验证；本轮研究模型未替换在线模型，也不能替代正式检测报告。')
    doc.add_heading('数据来源与容量代理定义', 1)
    paragraph(doc, '数据来自公开的 20 辆 BAIC EU500 充电记录，原作者说明电池为 CATL NCM、额定容量 145 Ah。Iontech 为检索索引，实际下载仓库与作者仓库分别列于参考文献。本轮完整读取 20 个原始 CSV 并重新提取片段，与 V2 特征和统计进行核验后，使用最终公开匿名特征表重新训练；源文件摘要与逐条预测保留追溯关系。')
    paragraph(doc, '容量代理定义为 Q = −∫I(t)dt / 3600 / (ΔSOC/100)，单位 Ah。提取时剔除重复时间、无效数值及异常范围；在原始行序建立无效观测屏障后排序，时间缺口大于 10 秒时切段。保留持续至少 1200 秒、SOC 增量至少 20 个百分点、SOC 单调且电流全负的片段，并筛选 0 < Q ≤ 250 Ah。250 Ah 为研究筛选界限，不是安全阈值。')
    paragraph(doc, '输入限定为前 600 秒起始 SOC、电压轨迹摘要、平均电流、温度摘要和单体压差摘要共 10 项特征。未将车辆 ID、全程时长、全程 SOC 增量或可用容量列输入模型。前缀电流与 SOC 仍与代理标签共享数学和物理信息，不能据此宣称标签独立。')
    doc.add_page_break()
    doc.add_heading('车辆隔离与固定实验设计', 1)
    figure_doc(doc, output, 1, notes[0])
    table_doc(doc, '表 1  固定实验设置', ['项目', '设置'], [
        ['划分', '训练车 1—14；验证车 15—16；测试车 17—20'],
        ['基线', '训练均值；标准化 Ridge（α=10）；随机森林'],
        ['随机森林', '200 棵树；叶节点最少 5 条；主种子 42'],
        ['重复与分组', '种子 17、42、2026；训练车内五折 GroupKFold'],
        ['消融', '分别去 SOC/电流、温度、单体压差后重新拟合'],
        ['压力测试', '训练标准差比例高斯噪声 1%/5%；随机缺失 10% 并用训练中位数填补']])
    paragraph(doc, '全部参数固定，不根据测试结果选择最优模型或种子。V2 的补充方案在已知早期测试结果之后制定，V3 为最终复核与重现，不是事先注册的独立盲测；验证集与固定测试集均保留原有角色。')
    doc.add_page_break()
    doc.add_heading('固定测试集结果', 1)
    figure_doc(doc, output, 2, notes[1], 6.0)
    table_doc(doc, '表 2  三模型固定测试集指标', ['模型', 'MAE / Ah', 'RMSE / Ah', 'R²'], [
        [LABELS[r.experiment], f'{r.mae_ah:.4f}', f'{r.rmse_ah:.4f}', f'{r.r2:.4f}'] for _, r in data['base'].iterrows()])
    paragraph(doc, f'测试集含 {int(rf.n):,} 段、{int(rf.vehicles)} 辆车。随机森林 RMSE 为 {rf.rmse_ah:.4f} Ah，R² 为 {rf.r2:.4f}。全部对应逐条预测均随研究产物保存，可用于独立指标复算。')
    paragraph(doc, notes[3])
    doc.add_heading('解释限制', 2)
    paragraph(doc, '测试车仅四辆，且属于同车型和同一公开数据源。连续充电片段相关；结果不能直接推广到其他车型、化学体系、全生命周期工况或真实退役决策。片段级平均误差与车辆等权平均误差分别回答不同问题。')
    doc.add_page_break()
    doc.add_heading('稳健性与特征依赖', 1)
    figure_doc(doc, output, 5, notes[4], 5.8, display_number=3)
    figure_doc(doc, output, 7, notes[6], 5.8, display_number=4)
    doc.add_page_break()
    doc.add_heading('压力测试与质量复核', 1)
    figure_doc(doc, output, 8, notes[7], 5.8, display_number=5)
    table_doc(doc, '表 3  质量复核状态', ['检查项目', '状态'], [[r.label, r.status] for _, r in data['quality'].iterrows()])
    paragraph(doc, notes[9])
    doc.add_page_break()
    doc.add_heading('研究边界与下一步验证', 1)
    paragraph(doc, '容量代理依赖 SOC 估计与电流测量，其误差可能来自 SOC 标定、积分偏差、温度及充电策略。特征消融反映模型在当前数据上的依赖，不证明因果机制；模拟特征噪声也不代表传感器标定后的真实噪声模型。')
    paragraph(doc, '下一步需要获得独立容量测试及车辆工况标签，开展跨车型、跨地区和前瞻性验证；寿命研究需要失效定义、随访及右删失处理；循环利用决策需要真实检测、安全门槛、经济与碳核算边界及处置结局。本轮不将这些待完成工作写成已取得成果。')
    paragraph(doc, '本轮未使用 IVST 长安或 Oxford 数据开展训练。公开实车容量代理实验与已有实验室电芯验证是不同证据域，不相互替代。')
    references(doc)
    doc.save(docs / '03_电循智策_挑战杯研究证据报告_最终复核版.docx')

    doc = new_doc()
    doc.add_paragraph(TITLE, 'Title')
    doc.add_paragraph('模型实验结果及图表  最终复核版 V3', 'Subtitle')
    paragraph(doc, '本册集中提供十幅可追溯学术图表。图内只呈现必要坐标、单位、统计说明；完整图题与解释置于图下。PNG 为 300 dpi，另有 SVG 与逐图 CSV 源数据。所有指标由本轮结果文件生成。')
    doc.add_heading('阅读约定', 1)
    paragraph(doc, '除训练车辆内五折验证图外，模型误差均来自固定测试车 17—20。标签为由电流积分和 SOC 变化得到的容量代理，单位 Ah。MAE、RMSE 越低越好；R² 为决定系数，不是准确率百分比。')
    paragraph(doc, '随机种子与五折误差线均使用样本标准差（ddof=1），不作为置信区间。压力测试仅改变输入特征，保持标签和训练完成的模型不变。十幅图共同展示研究证据，不构成实车健康、安全或退役处置认证。')
    table_doc(doc, '表 1  图表索引', ['图号', '主题'], [[i+1, c] for i, c in enumerate(CAPTIONS)])
    for start in range(1, 11, 2):
        doc.add_page_break()
        for number in [start, start+1]: figure_doc(doc, output, number, notes[number-1], 5.85)
    doc.add_page_break(); references(doc)
    doc.save(docs / '04_电循智策_模型实验结果及图表_最终复核版.docx')


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--result', type=Path, default=Path('research/public_battery/results/v3'))
    p.add_argument('--source', type=Path, default=Path('research/public_battery/data/v3/session_features.csv'))
    p.add_argument('--output', type=Path, required=True)
    p.add_argument('--verification', type=Path, required=True)
    p.add_argument('--figures-only', action='store_true')
    args = p.parse_args()
    data = build_figures(args.result, args.source, args.output, args.verification)
    if not args.figures_only: build_documents(data, args.output)
    notes = interpretations(data)
    report = '# ' + TITLE + '\n\n模型实验结果及图表 V3\n\n'
    report += '\n\n'.join(f'## 图 {i+1} {caption}\n\n{notes[i]}' for i, caption in enumerate(CAPTIONS))
    (args.output / '09_reports' / 'figure_interpretation.md').write_text(report, encoding='utf-8')
    print(json.dumps({'figures': 10, 'documents': 0 if args.figures_only else 2, 'all_checks_passed': bool(data['quality'].passed.all())}))


if __name__ == '__main__': main()
